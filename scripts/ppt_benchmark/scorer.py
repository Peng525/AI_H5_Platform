"""Score exported benchmark PPTX files."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

from .cost import DEEPSEEK_PROFILE, DEEPSEEK_USAGE_URL
from .prompts import REQUIRED_KEYWORDS


def extract_slide_texts(pptx_path: Path) -> list[str]:
    prs = Presentation(str(pptx_path))
    texts: list[str] = []
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        parts.append(t)
        texts.append("\n".join(parts))
    return texts


def count_editable_text_shapes(pptx_path: Path) -> tuple[int, int]:
    prs = Presentation(str(pptx_path))
    text_shapes = 0
    total = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            total += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_shapes += 1
    return text_shapes, total


def score_structure(page_count: int, target: int = 8) -> tuple[int, str]:
    if page_count == target:
        return 5, f"严格 {target} 页"
    if abs(page_count - target) == 1:
        return 3, f"{page_count} 页（差 1 页）"
    return 1, f"{page_count} 页，结构偏离"


def score_keywords(all_text: str, bench_type: str) -> tuple[int, str]:
    keys = REQUIRED_KEYWORDS.get(bench_type, [])
    if not keys:
        return 3, "无关键词表"
    hit = sum(1 for k in keys if k in all_text)
    ratio = hit / len(keys)
    if ratio >= 0.9:
        return 5, f"关键数据 {hit}/{len(keys)}"
    if ratio >= 0.6:
        return 3, f"部分数据 {hit}/{len(keys)}"
    return 1, f"数据缺失 {hit}/{len(keys)}"


def score_editability(text_shapes: int, total_shapes: int) -> tuple[int, str]:
    if text_shapes >= 8:
        return 5, f"可编辑文本形状 {text_shapes}"
    if text_shapes >= 4:
        return 3, f"部分可编辑 {text_shapes}/{total_shapes}"
    return 1, f"几乎无文本 {text_shapes}/{total_shapes}"


def score_layout_heuristic(all_text: str) -> tuple[int, str]:
    # 无视觉检测时的保守启发式
    if len(all_text) < 80:
        return 2, "文本过少，可能版式空或图化"
    if len(all_text) > 12000:
        return 2, "文本过多，可能溢出"
    return 4, "文本量正常（未做像素级检测）"


def score_style_heuristic(slide_texts: list[str]) -> tuple[int, str]:
    if len(slide_texts) < 2:
        return 1, "页数不足"
    lengths = [len(t) for t in slide_texts]
    avg = sum(lengths) / len(lengths)
    var = sum((l - avg) ** 2 for l in lengths) / len(lengths)
    if var < avg * avg * 4:
        return 4, "各页信息量较均衡"
    return 3, "各页信息量差异较大"


def score_content_llm_placeholder(bench_type: str, all_text: str) -> tuple[int, str]:
    # 自动化流水线：基于关键词+长度；完整文案质量需人工复核
    kw_score, kw_note = score_keywords(all_text, bench_type)
    if kw_score >= 4 and len(all_text) > 200:
        return 4, f"自动评估：数据覆盖好（{kw_note}）"
    return kw_score, f"自动评估：{kw_note}"


def score_file(pptx_path: Path, bench_type: str, elapsed_min: float, est_usd: float) -> dict:
    slide_texts = extract_slide_texts(pptx_path)
    all_text = "\n".join(slide_texts)
    text_shapes, total_shapes = count_editable_text_shapes(pptx_path)

    s1, n1 = score_structure(len(slide_texts))
    s2, n2 = score_layout_heuristic(all_text)
    s3, n3 = score_content_llm_placeholder(bench_type, all_text)
    s4, n4 = score_keywords(all_text, bench_type)
    s5, n5 = score_style_heuristic(slide_texts)
    s6, n6 = score_editability(text_shapes, total_shapes)

    if elapsed_min <= 15:
        s7, n7 = 5, f"{elapsed_min:.1f} min"
    elif elapsed_min <= 30:
        s7, n7 = 3, f"{elapsed_min:.1f} min"
    else:
        s7, n7 = 1, f"{elapsed_min:.1f} min"

    if est_usd <= 0.5:
        s8, n8 = 5, f"${est_usd:.3f}"
    elif est_usd <= 2.0:
        s8, n8 = 3, f"${est_usd:.3f}"
    else:
        s8, n8 = 1, f"${est_usd:.3f}"

    core = (s1 + s2 + s3 + s4 + s5 + s6) / 6
    return {
        "page_count": len(slide_texts),
        "scores": {
            "结构完整度": (s1, n1),
            "版式精度": (s2, n2),
            "文案质量": (s3, n3),
            "图表数据呈现": (s4, n4),
            "风格一致性": (s5, n5),
            "可编辑性": (s6, n6),
            "耗时": (s7, n7),
            "预估费用": (s8, n8),
        },
        "core_avg": round(core, 2),
        "all_text_preview": all_text[:500],
    }


def format_result_md(
    *,
    date_str: str,
    profile: str,
    bench_type: str,
    project_path: Path,
    export_path: Path,
    elapsed_min: float,
    est_usd: int | float,
    score: dict,
    phase: int = 1,
    error: str = "",
    cost: dict | None = None,
    usage: dict | None = None,
    page_records: list[dict] | None = None,
    cost_json: Path | None = None,
) -> str:
    lines = [
        f"# 横评记录 {date_str}-{profile}-{bench_type}-phase{phase}",
        "",
        "## 基本信息",
        "",
        "| 字段 | 值 |",
        "|------|-----|",
        f"| 日期 | {date_str} |",
        f"| active_profile | {profile} |",
        f"| benchmark_phase | {phase} |",
        f"| 类型 | {bench_type} |",
        f"| ppt-master 项目路径 | {project_path} |",
        f"| exports 路径 | {export_path} |",
        f"| 总耗时（分钟） | {elapsed_min:.1f} |",
        f"| 实际页数 | {score.get('page_count', 0)} |",
        "",
    ]
    if error:
        lines.extend([f"**错误**: {error}", ""])

    lines.extend(["## Phase 1 评分（1–5）", "", "| 维度 | 分 | 备注 |", "|------|-----|------|"])
    for dim, (val, note) in score.get("scores", {}).items():
        if dim in ("耗时", "预估费用"):
            continue
        lines.append(f"| {dim} | {val} | {note} |")
    lines.append(f"| **核心均分 (1–6)** | **{score.get('core_avg', '-')}** | |")
    lines.extend(
        [
            "",
            "## 成本与问题",
            "",
            "| 字段 | 值 |",
            "|------|-----|",
        ]
    )
    if cost and usage and profile == DEEPSEEK_PROFILE:
        lines.extend(
            [
                f"| prompt_tokens | {usage.get('prompt_tokens', 0):,} |",
                f"| completion_tokens | {usage.get('completion_tokens', 0):,} |",
                f"| total_tokens | {usage.get('total_tokens', 0):,} |",
            ]
        )
        if usage.get("prompt_cache_hit_tokens") or usage.get("prompt_cache_miss_tokens"):
            lines.append(
                f"| prompt_cache_hit / miss | "
                f"{usage.get('prompt_cache_hit_tokens', 0):,} / "
                f"{usage.get('prompt_cache_miss_tokens', 0):,} |"
            )
        source = cost.get("pricing_source", "official")
        lines.extend(
            [
                f"| 定价来源 | {source} |",
                f"| estimated_usd_input | ${cost.get('estimated_usd_input', 0):.6f} |",
                f"| estimated_usd_output | ${cost.get('estimated_usd_output', 0):.6f} |",
                f"| estimated_usd_total | ${cost.get('estimated_usd_total', est_usd):.6f} |",
                f"| 账单对照 | [{DEEPSEEK_USAGE_URL}]({DEEPSEEK_USAGE_URL}) |",
            ]
        )
        if cost_json:
            lines.append(f"| cost JSON | {cost_json} |")
    elif profile == DEEPSEEK_PROFILE and est_usd:
        lines.append(f"| 预估 API 费用（USD） | ${est_usd:.4f} |")
    elif profile != DEEPSEEK_PROFILE:
        lines.append("| API 费用 | 自行查看中转账单 |")
    lines.append(f"| 耗时评分 | {score.get('scores', {}).get('耗时', ('', ''))[0]} |")

    if page_records and profile == DEEPSEEK_PROFILE:
        lines.extend(
            [
                "",
                "### 逐页 token / 费用",
                "",
                "| 页 | 文件 | prompt | completion | total | USD |",
                "|----|------|--------|------------|-------|-----|",
            ]
        )
        for pr in page_records:
            u = pr.get("usage") or {}
            lines.append(
                f"| {pr.get('page', '-')} | {pr.get('file', '-')} | "
                f"{u.get('prompt_tokens', 0):,} | {u.get('completion_tokens', 0):,} | "
                f"{u.get('total_tokens', 0):,} | "
                f"${pr.get('estimated_usd_total', 0):.6f} |"
            )

    lines.extend(
        [
            "",
            "---",
            "*由 develop/scripts/ppt_benchmark/run_phase1.py 自动生成*",
        ]
    )
    return "\n".join(lines)
