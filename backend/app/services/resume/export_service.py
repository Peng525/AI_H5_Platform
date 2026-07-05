"""Export resume to PDF/DOCX with template-aware layout."""
from __future__ import annotations

import io
from typing import Any

from app.services.resume.template_catalog import normalize_template_id
from app.services.resume.visual_compiler import normalize_structured

ACCENT = (37, 99, 235)
TEXT = (30, 41, 59)
MUTED = (100, 116, 139)

_TEMPLATE_ACCENTS: dict[str, tuple[int, int, int]] = {
    "template1": (37, 99, 235),
    "template2": (15, 118, 110),
    "template3": (79, 70, 229),
    "template4": (17, 24, 39),
    "template5": (30, 64, 175),
    "template6": (190, 24, 93),
    "template7": (22, 163, 74),
}


def _template_accent(visual: dict[str, Any] | None) -> tuple[int, int, int]:
    tid = normalize_template_id((visual or {}).get("template_id")) or "template1"
    return _TEMPLATE_ACCENTS.get(tid, ACCENT)


def _val(structured: dict[str, Any], *path: str, default: str = "") -> str:
    cur: Any = structured
    for p in path:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(p)
    return str(cur or default) if cur is not None else default


def _style_for(visual: dict[str, Any] | None, bind: str) -> dict[str, Any]:
    if not visual or not isinstance(visual.get("styles"), dict):
        return {}
    s = visual["styles"].get(bind)
    return s if isinstance(s, dict) else {}


def export_docx_bytes(structured: dict[str, Any], visual: dict[str, Any] | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    data = normalize_structured(structured)
    doc = Document()
    title = doc.add_heading("个人简历", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    motto = _val(data, "basics", "motto")
    if motto:
        p = doc.add_paragraph(motto)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def section_heading(text: str) -> None:
        doc.add_heading(text, level=1)

    section_heading("基本信息")
    basics = data.get("basics") or {}
    for label, key in [
        ("姓名", "name"),
        ("电话", "phone"),
        ("邮箱", "email"),
        ("性别", "gender"),
        ("年龄", "age"),
    ]:
        v = basics.get(key, "")
        if v:
            doc.add_paragraph(f"{label}：{v}")

    ji = data.get("job_intention") or {}
    if any(ji.values()):
        section_heading("求职意向")
        for label, key in [("岗位", "position"), ("城市", "city"), ("薪资", "salary"), ("到岗", "availability")]:
            if ji.get(key):
                doc.add_paragraph(f"{label}：{ji[key]}")

    edu = data.get("education") or []
    if edu:
        section_heading("教育背景")
        for item in edu:
            if not isinstance(item, dict):
                continue
            line = " | ".join(filter(None, [item.get("period"), item.get("school"), item.get("degree")]))
            if line:
                doc.add_paragraph(line, style="List Bullet")
            if item.get("details"):
                doc.add_paragraph(str(item["details"]))

    exp = data.get("experience") or []
    if exp:
        section_heading("工作经历")
        for item in exp:
            if not isinstance(item, dict):
                continue
            header = " | ".join(filter(None, [item.get("period"), item.get("company"), item.get("title")]))
            if header:
                doc.add_paragraph(header)
            for b in item.get("bullets") or []:
                doc.add_paragraph(str(b), style="List Bullet")

    bars = data.get("skill_bars") or []
    if bars:
        section_heading("技能特长")
        for bar in bars:
            if isinstance(bar, dict):
                doc.add_paragraph(f"{bar.get('name', '')} — {bar.get('label', '')}")

    honors = data.get("honors") or []
    if honors:
        section_heading("荣誉证书")
        for h in honors:
            doc.add_paragraph(str(h), style="List Bullet")

    se = data.get("self_evaluation") or basics.get("summary")
    if se:
        section_heading("自我评价")
        doc.add_paragraph(str(se))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register_cjk_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "C:/Windows/Fonts/msyh.ttc",
    ]
    for path in candidates:
        try:
            pdfmetrics.registerFont(TTFont("ResumeCJK", path))
            return "ResumeCJK"
        except Exception:
            continue
    return "Helvetica"


def export_pdf_bytes(structured: dict[str, Any], visual: dict[str, Any] | None = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    data = normalize_structured(structured)
    visual = visual or {}
    accent = _template_accent(visual)
    tid = normalize_template_id(visual.get("template_id")) or "template1"
    basics = data.get("basics") or {}
    font = _register_cjk_font()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    x_margin = 40 if tid != "template5" else 160
    sidebar_x = 40 if tid == "template5" else None
    y = height - 45
    content_w = width - x_margin - 40

    def draw_section(title: str) -> None:
        nonlocal y
        if y < 80:
            c.showPage()
            y = height - 45
        if tid == "template4":
            c.setFillColorRGB(accent[0] / 255, accent[1] / 255, accent[2] / 255)
            c.setFont(font, 11)
            c.drawString(x_margin, y - 12, title)
            c.line(x_margin, y - 16, x_margin + content_w, y - 16)
            y -= 28
            c.setFillColorRGB(TEXT[0] / 255, TEXT[1] / 255, TEXT[2] / 255)
            return
        c.setFillColorRGB(accent[0] / 255, accent[1] / 255, accent[2] / 255)
        c.rect(x_margin if sidebar_x is None else sidebar_x, y - 18, content_w if sidebar_x is None else 100, 20, fill=1, stroke=0)
        c.setFillColorRGB(1, 1, 1)
        c.setFont(font, 11)
        c.drawString((x_margin if sidebar_x is None else sidebar_x) + 8, y - 14, title)
        y -= 32
        c.setFillColorRGB(TEXT[0] / 255, TEXT[1] / 255, TEXT[2] / 255)

    def draw_line(text: str, bind: str = "", size: int = 10) -> None:
        nonlocal y
        if not text:
            return
        if y < 50:
            c.showPage()
            y = height - 45
        st = _style_for(visual, bind)
        fs = int(st.get("fontSize") or size)
        c.setFont(font, fs)
        c.drawString(x_margin, y, text[:90])
        y -= fs + 6

    c.setFont(font, 18 if tid != "template6" else 22)
    title_text = str(basics.get("name") or "个人简历") if tid in ("template6", "template7") else "个人简历"
    c.drawCentredString(width / 2, y, title_text[:20])
    y -= 22
    motto = _val(data, "basics", "motto")
    if motto:
        c.setFont(font, 9)
        c.setFillColorRGB(MUTED[0] / 255, MUTED[1] / 255, MUTED[2] / 255)
        c.drawCentredString(width / 2, y, motto[:60])
        y -= 24
        c.setFillColorRGB(TEXT[0] / 255, TEXT[1] / 255, TEXT[2] / 255)

    draw_section("基本信息")
    row1 = f"姓名：{basics.get('name', '')}    电话：{basics.get('phone', '')}"
    row2 = f"邮箱：{basics.get('email', '')}    性别：{basics.get('gender', '')}    年龄：{basics.get('age', '')}"
    draw_line(row1, "basics.name")
    draw_line(row2, "basics.email")

    ji = data.get("job_intention") or {}
    if any(ji.values()):
        draw_section("求职意向")
        draw_line(
            f"岗位：{ji.get('position', '')}  城市：{ji.get('city', '')}  "
            f"薪资：{ji.get('salary', '')}  到岗：{ji.get('availability', '')}"
        )

    edu = data.get("education") or []
    if edu:
        draw_section("教育背景")
        for i, item in enumerate(edu):
            if not isinstance(item, dict):
                continue
            draw_line(
                f"{item.get('period', '')}    {item.get('school', '')}    {item.get('degree', '')}",
                f"education[{i}].school",
            )
            if item.get("details"):
                draw_line(f"  {item['details']}", f"education[{i}].details", 9)

    exp = data.get("experience") or []
    if exp:
        draw_section("工作经历")
        for i, item in enumerate(exp):
            if not isinstance(item, dict):
                continue
            draw_line(
                f"{item.get('period', '')}    {item.get('company', '')}    {item.get('title', '')}",
                f"experience[{i}].company",
            )
            for j, b in enumerate(item.get("bullets") or []):
                draw_line(f"  • {b}", f"experience[{i}].bullets[{j}]", 9)

    bars = data.get("skill_bars") or []
    if bars:
        draw_section("技能特长")
        for i, bar in enumerate(bars):
            if isinstance(bar, dict):
                draw_line(f"{bar.get('name', '')} — {bar.get('label', '')}", f"skill_bars[{i}].name")

    honors = data.get("honors") or []
    if honors:
        draw_section("荣誉证书")
        for i, h in enumerate(honors):
            draw_line(f"• {h}", f"honors[{i}]")

    se = data.get("self_evaluation") or basics.get("summary")
    if se:
        draw_section("自我评价")
        draw_line(str(se), "self_evaluation")

    c.save()
    return buf.getvalue()


def structured_to_plain_text(structured: dict[str, Any]) -> str:
    """Legacy plain export helper."""
    data = normalize_structured(structured)
    lines: list[str] = []
    basics = data.get("basics") or {}
    if basics.get("name"):
        lines.append(str(basics["name"]))
    contact = " · ".join(filter(None, [basics.get("phone"), basics.get("email")]))
    if contact:
        lines.append(contact)
    for item in data.get("experience") or []:
        if isinstance(item, dict):
            lines.append(str(item.get("company", "")))
    return "\n".join(lines).strip()
