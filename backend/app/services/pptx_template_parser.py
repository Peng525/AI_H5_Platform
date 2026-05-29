"""将 PPTX 解析为 H5 模板 slides_json 结构。"""
from __future__ import annotations

import base64
import io
import re
import uuid
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

DEFAULT_SETTINGS = {
    "viewportId": "mobile-375",
    "scrollEffect": "page",
    "themeId": "zjy-minimal",
    "showScrollHint": False,
    "defaultChatTapToContinue": True,
    "bgm": {"enabled": False, "trackId": "", "url": "", "loop": True, "volume": 0.35},
}


class PptxParseError(Exception):
    pass


def _slugify(name: str) -> str:
    base = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", name.strip().lower())
    base = re.sub(r"-+", "-", base).strip("-")
    if not base:
        base = "imported-template"
    return f"{base[:48]}-{uuid.uuid4().hex[:6]}"


def _emu_to_px(emu: int, slide_emu: int, canvas_px: int) -> int:
    if not slide_emu:
        return 0
    return max(0, round((emu / slide_emu) * canvas_px))


def _rgb_to_hex(color) -> str | None:
    if color is None:
        return None
    try:
        if color.type == MSO_COLOR_TYPE.RGB and color.rgb is not None:
            rgb = color.rgb
            return f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    except (TypeError, ValueError, AttributeError):
        return None
    return None


def _shape_fill_color(shape) -> str | None:
    try:
        fill = shape.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            return _rgb_to_hex(fill.fore_color)
    except (AttributeError, ValueError, TypeError):
        return None
    return None


def _text_style(shape) -> dict[str, Any]:
    style: dict[str, Any] = {
        "background": "transparent",
        "textAlign": "left",
        "fontSize": 16,
        "fontWeight": "normal",
        "color": "#1B1B1C",
        "fontFamily": '"Microsoft YaHei", "微软雅黑", sans-serif',
        "lineHeight": 1.4,
    }
    try:
        if not shape.has_text_frame:
            return style
        para = shape.text_frame.paragraphs[0]
        if para.runs:
            run = para.runs[0]
            if run.font.size:
                style["fontSize"] = max(10, min(72, int(run.font.size.pt)))
            if run.font.bold:
                style["fontWeight"] = "bold"
            color = _rgb_to_hex(run.font.color) if run.font.color else None
            if color:
                style["color"] = color
        align = str(para.alignment).lower() if para.alignment is not None else ""
        if "center" in align:
            style["textAlign"] = "center"
        elif "right" in align:
            style["textAlign"] = "right"
    except (AttributeError, ValueError, TypeError, IndexError):
        pass
    return style


def _slide_background(slide) -> str:
    try:
        fill = slide.background.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            color = _rgb_to_hex(fill.fore_color)
            if color:
                return color
    except (AttributeError, ValueError, TypeError):
        pass
    return "#FFFFFF"


def _parse_shape(
    shape,
    slide_idx: int,
    z: int,
    ppt_w: int,
    ppt_h: int,
    canvas_w: int,
    canvas_h: int,
) -> tuple[dict[str, Any] | None, int]:
    if not hasattr(shape, "width"):
        return None, z

    x = _emu_to_px(shape.left, ppt_w, canvas_w)
    y = _emu_to_px(shape.top, ppt_h, canvas_h)
    w = max(8, _emu_to_px(shape.width, ppt_w, canvas_w))
    h = max(8, _emu_to_px(shape.height, ppt_h, canvas_h))

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            blob = shape.image.blob
            content_type = shape.image.content_type or "image/png"
            data_url = f"data:{content_type};base64,{base64.b64encode(blob).decode('ascii')}"
            return {
                "id": f"img_{slide_idx}_{z}",
                "type": "image",
                "x": x,
                "y": y,
                "width": w,
                "height": h,
                "zIndex": z,
                "content": data_url,
                "style": {"background": "#F0F0F0", "objectFit": "cover"},
            }, z + 1
        except (AttributeError, ValueError, TypeError):
            return None, z

    if shape.has_text_frame:
        text = (shape.text_frame.text or "").strip()
        if not text:
            fill = _shape_fill_color(shape)
            if fill:
                return {
                    "id": f"s_{slide_idx}_{z}",
                    "type": "shape",
                    "x": x,
                    "y": y,
                    "width": w,
                    "height": h,
                    "zIndex": z,
                    "content": "",
                    "style": {"background": fill, "borderRadius": 4},
                }, z + 1
            return None, z
        return {
            "id": f"t_{slide_idx}_{z}",
            "type": "text",
            "x": x,
            "y": y,
            "width": w,
            "height": max(h, 24),
            "zIndex": z,
            "content": text,
            "style": _text_style(shape),
        }, z + 1

    fill = _shape_fill_color(shape)
    if fill:
        return {
            "id": f"s_{slide_idx}_{z}",
            "type": "shape",
            "x": x,
            "y": y,
            "width": w,
            "height": h,
            "zIndex": z,
            "content": "",
            "style": {"background": fill, "borderRadius": 4},
        }, z + 1

    return None, z


def parse_pptx_bytes(
    file_bytes: bytes,
    *,
    device: str = "mobile",
    title: str | None = None,
    category: str = "简约商务",
) -> dict[str, Any]:
    if not file_bytes:
        raise PptxParseError("文件为空")
    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise PptxParseError(f"无法读取 PPTX：{exc}") from exc

    if not prs.slides:
        raise PptxParseError("PPT 中没有幻灯片")

    canvas_w = 375 if device == "mobile" else 1280
    canvas_h = 812 if device == "mobile" else 720
    ppt_w = int(prs.slide_width or 1)
    ppt_h = int(prs.slide_height or 1)
    default_viewport = "mobile-375" if device == "mobile" else "web-1280"

    slides: list[dict[str, Any]] = []
    for slide_idx, slide in enumerate(prs.slides):
        elements: list[dict[str, Any]] = []
        z = 1
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            el, z = _parse_shape(shape, slide_idx, z, ppt_w, ppt_h, canvas_w, canvas_h)
            if el:
                elements.append(el)

        slide_title = ""
        for el in elements:
            if el["type"] == "text" and el.get("content"):
                slide_title = str(el["content"]).split("\n")[0][:40]
                break
        if not slide_title:
            slide_title = f"第 {slide_idx + 1} 页"

        slides.append(
            {
                "title": slide_title,
                "subtitle": "",
                "layout": "cover" if slide_idx == 0 else "bullets",
                "animation": "fade",
                "canvas_background": _slide_background(slide),
                "canvas_elements": elements,
                "bullets": [],
            }
        )

    inferred_title = title or "导入的演示模板"
    settings = {
        **DEFAULT_SETTINGS,
        "viewportId": default_viewport,
    }
    return {
        "id": _slugify(inferred_title),
        "title": inferred_title,
        "description": f"由 PPT 自动解析，共 {len(slides)} 页。",
        "category": category,
        "device": device,
        "pages": len(slides),
        "premium": False,
        "cover_gradient": "from-primary to-primary-container",
        "default_viewport": default_viewport,
        "slides_json": slides,
        "settings_json": settings,
        "enabled": True,
        "sort_order": 0,
    }
